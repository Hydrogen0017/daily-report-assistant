# -*- coding: utf-8 -*-
"""
Flask 后端服务器 - 提供 API 路由
处理前端请求，调用数据处理模块
"""
import os
import sys
import json
import shutil
import threading
from pathlib import Path
from flask import Flask, request, jsonify, send_file, send_from_directory

# 确保项目根目录在 path 中
BASE_DIR = Path(__file__).parent
sys.path.insert(0, str(BASE_DIR))

from store import (logger, load_config, save_config, get_today_str, get_now_str,
                   UPLOADS_DIR, OUTPUT_DIR, CACHE_DIR, LOGS_DIR, TEMPLATES_DIR, DATA_DIR)
from processor import (DailyReportProcessor, update_light_oil_sales,
                       update_monthly_targets, update_station_regions,
                       update_quan_coefficients, UPLOAD_FILES)

app = Flask(__name__, static_folder=str(BASE_DIR / "frontend"), static_url_path="")

# 全局处理器实例（每个会话独立）
processors = {}
# 处理状态
processing_status = {"daily_report": {"status": "idle", "progress": 0, "message": "", "steps": []}}
# 锁
process_lock = threading.Lock()


# ===== 静态文件 =====
@app.route("/")
def index():
    return send_from_directory(str(BASE_DIR / "frontend"), "index.html")


# ===== API: 应用信息 =====
@app.route("/api/app/info")
def app_info():
    version_file = BASE_DIR / "version.json"
    with open(version_file, "r", encoding="utf-8") as f:
        version_info = json.load(f)
    config = load_config()
    return jsonify({
        "name": version_info["name"],
        "version": version_info["version"],
        "reports": version_info["reports"],
        "report_date": get_today_str(),
        "last_run": config.get("last_run"),
        "settings": config.get("settings", {})
    })


# ===== API: 报表入口列表 =====
@app.route("/api/reports/list")
def reports_list():
    version_file = BASE_DIR / "version.json"
    with open(version_file, "r", encoding="utf-8") as f:
        version_info = json.load(f)
    return jsonify({"reports": version_info["reports"]})


# ===== API: 非油日报 - 上传文件 =====
@app.route("/api/daily-report/upload", methods=["POST"])
def upload_files():
    """上传6张导出表"""
    session_id = f"session_{get_today_str()}_{int(threading.get_ident())}"
    session_dir = UPLOADS_DIR / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    uploaded = {}
    errors = []
    for key, expected_name in UPLOAD_FILES.items():
        if key in request.files:
            f = request.files[key]
            save_path = session_dir / f.filename
            f.save(str(save_path))
            uploaded[key] = str(save_path)
            logger.info(f"上传 {expected_name}: {f.filename}")
        else:
            errors.append({"key": key, "name": expected_name, "error": "未上传"})

    # 保存会话信息
    session_info = {"session_id": session_id, "files": uploaded, "time": get_now_str()}
    session_file = CACHE_DIR / f"{session_id}.json"
    with open(session_file, "w", encoding="utf-8") as f:
        json.dump(session_info, f, ensure_ascii=False, indent=2)

    return jsonify({
        "status": "success" if len(uploaded) == 6 else "partial",
        "session_id": session_id,
        "uploaded": uploaded,
        "errors": errors,
        "message": f"已上传 {len(uploaded)}/6 个文件"
    })


# ===== API: 非油日报 - 多文件自动识别上传 =====
# 文件名 → 内部 key 的关键词映射
FILE_NAME_KEYWORDS = {
    "dangqi_all": ["当期全", "当期全部"],
    "tongqi_all": ["同期全", "同期全部"],
    "dangqi_ling": ["当期零", "当期零售"],
    "crm": ["CRM", "crm"],
    "dianzi_quan": ["电子券", "非油券"],
    "yangche_ka": ["养车卡", "养车"],
}


def identify_file_key(filename):
    """根据文件名自动识别文件类型"""
    name = str(filename)
    for key, keywords in FILE_NAME_KEYWORDS.items():
        for kw in keywords:
            if kw in name:
                return key
    return None


@app.route("/api/daily-report/upload-multi", methods=["POST"])
def upload_multi_files():
    """多文件上传，自动识别文件类型（一次选择6个文件）"""
    session_id = f"session_{get_today_str()}_{int(threading.get_ident())}"
    session_dir = UPLOADS_DIR / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    files = request.files.getlist("files")
    if not files:
        return jsonify({"status": "error", "message": "未收到文件"}), 400

    uploaded = {}
    matched = {}
    unmatched = []

    for f in files:
        fname = f.filename
        key = identify_file_key(fname)
        if key:
            save_path = session_dir / fname
            f.save(str(save_path))
            uploaded[key] = str(save_path)
            matched[key] = fname
            logger.info(f"自动识别 {fname} -> {key} ({UPLOAD_FILES.get(key, '?')})")
        else:
            unmatched.append(fname)
            logger.warn(f"无法识别文件类型: {fname}")

    # 保存会话信息
    session_info = {"session_id": session_id, "files": uploaded,
                    "matched": matched, "unmatched": unmatched, "time": get_now_str()}
    session_file = CACHE_DIR / f"{session_id}.json"
    with open(session_file, "w", encoding="utf-8") as f:
        json.dump(session_info, f, ensure_ascii=False, indent=2)

    # 返回时 uploaded 用文件名便于前端展示
    uploaded_names = {k: Path(v).name for k, v in uploaded.items()}

    return jsonify({
        "status": "success" if len(matched) == 6 else "partial",
        "session_id": session_id,
        "uploaded": uploaded_names,
        "matched": matched,
        "unmatched": unmatched,
        "message": f"已识别 {len(matched)}/6 个文件"
    })


# ===== API: 非油日报 - 获取上传文件列表 =====
@app.route("/api/daily-report/uploads/<session_id>")
def get_uploads(session_id):
    session_file = CACHE_DIR / f"{session_id}.json"
    if session_file.exists():
        with open(session_file, "r", encoding="utf-8") as f:
            return jsonify(json.load(f))
    return jsonify({"status": "error", "message": "会话不存在"}), 404


# ===== API: 非油日报 - 提交调整数据 =====
@app.route("/api/daily-report/adjustments", methods=["POST"])
def submit_adjustments():
    """提交手动调整数据"""
    data = request.json
    session_id = data.get("session_id")
    adjustments = data.get("adjustments", [])

    # 保存调整到会话缓存
    adj_file = CACHE_DIR / f"{session_id}_adjustments.json"
    with open(adj_file, "w", encoding="utf-8") as f:
        json.dump({"session_id": session_id, "adjustments": adjustments, "time": get_now_str()},
                  f, ensure_ascii=False, indent=2)
    logger.info(f"提交调整: {len(adjustments)}条")
    return jsonify({"status": "success", "count": len(adjustments)})


# ===== API: 非油日报 - 执行处理 =====
@app.route("/api/daily-report/process", methods=["POST"])
def process_report():
    """执行日报处理（后台线程）"""
    data = request.json or {}
    session_id = data.get("session_id")
    adjustments = data.get("adjustments", [])
    output_format = data.get("output_format", "both")  # xlsx / image / both

    if not session_id:
        return jsonify({"status": "error", "message": "缺少 session_id"}), 400

    session_file = CACHE_DIR / f"{session_id}.json"
    if not session_file.exists():
        return jsonify({"status": "error", "message": "会话不存在，请先上传文件"}), 400

    with open(session_file, "r", encoding="utf-8") as f:
        session_info = json.load(f)

    def progress_cb(info):
        processing_status["daily_report"] = {
            "status": info["status"],
            "progress": info["progress"],
            "message": info["message"],
            "step": info["step"],
            "total": info["total"],
            "time": get_now_str()
        }

    def run_process():
        try:
            processing_status["daily_report"] = {"status": "processing", "progress": 0,
                "message": "开始处理...", "step": 0, "total": 10, "time": get_now_str()}
            logger.start_session("daily_report")

            config = load_config()
            processor = DailyReportProcessor(config)
            processor.set_progress_callback(progress_cb)

            # 设置调整
            for adj in adjustments:
                processor.add_adjustment(adj["station"], adj["type"], adj["value"], adj.get("reason", ""))

            result = processor.run_all(session_info["files"], adjustments)

            output_paths = {}
            if result["status"] == "success":
                if output_format in ("xlsx", "both"):
                    template = data.get("template_path")
                    output_paths["xlsx"] = processor.output_xlsx(template)
                if output_format in ("image", "both"):
                    output_paths["image"] = processor.output_image()

                result["outputs"] = output_paths
                # 更新配置
                config["last_run"] = get_now_str()
                config["report_date"] = processor.report_date
                save_config(config)

            processing_status["daily_report"] = {
                "status": "done" if result["status"] == "success" else "error",
                "progress": 100,
                "message": result.get("message", "处理完成"),
                "result": result,
                "time": get_now_str()
            }
        except Exception as e:
            logger.error(f"处理异常: {str(e)}")
            processing_status["daily_report"] = {
                "status": "error", "progress": 100,
                "message": f"处理异常: {str(e)}", "time": get_now_str()
            }

    thread = threading.Thread(target=run_process, daemon=True)
    thread.start()
    return jsonify({"status": "started", "session_id": session_id, "message": "处理已启动"})


# ===== API: 非油日报 - 查询处理进度 =====
@app.route("/api/daily-report/status")
def process_status():
    return jsonify(processing_status.get("daily_report", {"status": "idle"}))


# ===== API: 非油日报 - 预览数据 =====
@app.route("/api/daily-report/preview")
def preview_data():
    """获取预览数据，返回结构化通报数据供前端按原日报格式渲染"""
    session_id = request.args.get("session_id")
    outputs = sorted(OUTPUT_DIR.glob("每日通报*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not outputs:
        return jsonify({"status": "error", "message": "暂无输出文件"})

    import openpyxl
    wb = openpyxl.load_workbook(outputs[0], data_only=True)

    # 读取通报表，解析为结构化数据（新78列结构，数据行 R7-R12）
    tongbao_rows = []
    if "通报" in wb.sheetnames:
        ws = wb["通报"]
        # 列: A(1)序号 B(2)单位 C(3)日均吨油 D(4)门零吨油销售额 E(5)门零销售
        #     G(7)目标 H(8)完成含 I(9)完成剔 J(10)完成率 K(11)同期 L(12)增减量 M(13)增幅
        #     O(15)毛利目标 P(16)毛利含 Q(17)毛利剔 R(18)毛利完成率 S(19)毛利率
        #     T(20)名次 V(22)综合完成率
        raw_rows = []
        for row in ws.iter_rows(min_row=7, max_row=12, values_only=True):
            raw_rows.append(list(row))

        def v(vals, c):
            return vals[c] if c < len(vals) else None

        for i, vals in enumerate(raw_rows):
            if len(vals) >= 23:
                if i == 0:
                    # 合计行是公式，data_only 读取为 None，手动计算
                    # 数值列求和: C(2) D(3) E(4) G(6) H(7) I(8) K(10) L(11) O(14) P(15) Q(16)
                    for c in [2, 3, 4, 6, 7, 8, 10, 11, 14, 15, 16]:
                        col_sum = 0
                        for j in range(1, len(raw_rows)):
                            x = v(raw_rows[j], c)
                            col_sum += x if isinstance(x, (int, float)) else 0
                        vals[c] = col_sum
                    # 率列计算
                    g, h_, ii = v(vals,6), v(vals,7), v(vals,8)
                    o, p_, q = v(vals,14), v(vals,15), v(vals,16)
                    k_ = v(vals,10)
                    vals[9] = (ii / g) if g else 0                 # J 完成率
                    vals[12] = ((v(vals,11) or 0) / k_) if k_ else None  # M 增幅
                    vals[17] = (q / o) if o else 0                 # R 毛利完成率
                    vals[18] = (p_ / h_) if h_ else 0              # S 毛利率
                    vals[21] = ((q / o) * 0.5 + min(ii / g, 1.3) * 0.4) if (o and g) else 0  # V 综合

                tongbao_rows.append({
                    "序号": v(vals,0), "单位": v(vals,1),
                    "日均吨油": v(vals,2), "门零吨油销售额": v(vals,3), "门零销售": v(vals,4),
                    "基础品类目标": v(vals,6), "基础品类完成含非非": v(vals,7),
                    "基础品类完成剔除非非": v(vals,8), "基础品类完成率": v(vals,9),
                    "同期": v(vals,10), "增减量": v(vals,11), "增幅": v(vals,12),
                    "毛利目标": v(vals,14), "毛利完成含非非": v(vals,15),
                    "毛利完成剔除非非": v(vals,16), "毛利完成率": v(vals,17),
                    "毛利率": v(vals,18), "名次": v(vals,19), "综合完成率": v(vals,21)
                })

    # 附表数据（前20行摘要）
    fuhe_data = []
    if "附表" in wb.sheetnames:
        ws = wb["附表"]
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True):
            fuhe_data.append(list(row))

    # 通报图片路径
    img_path = OUTPUT_DIR / f"每日通报{get_today_str()}.png"
    # 找最新的 png
    pngs = sorted(OUTPUT_DIR.glob("每日通报*.png"), key=lambda p: p.stat().st_mtime, reverse=True)

    return jsonify({
        "status": "success",
        "file": str(outputs[0]),
        "tongbao_rows": tongbao_rows,
        "fuhe": fuhe_data,
        "report_date": get_today_str(),
        "image_url": f"/api/download/{pngs[0].name}" if pngs else None
    })


# ===== API: 下载输出文件 =====
@app.route("/api/download/<path:filename>")
def download_file(filename):
    file_path = OUTPUT_DIR / filename
    if file_path.exists():
        return send_file(str(file_path), as_attachment=True)
    return jsonify({"status": "error", "message": "文件不存在"}), 404


# ===== API: 预览图片（非 attachment，直接显示） =====
@app.route("/api/serve/<path:filename>")
def serve_file(filename):
    file_path = OUTPUT_DIR / filename
    if file_path.exists():
        return send_file(str(file_path), as_attachment=False)
    return jsonify({"status": "error", "message": "文件不存在"}), 404


# ===== API: 用系统默认程序打开文件 =====
@app.route("/api/open/<path:filename>", methods=["POST"])
def open_file(filename):
    import subprocess, platform
    file_path = OUTPUT_DIR / filename
    if not file_path.exists():
        return jsonify({"status": "error", "message": "文件不存在"}), 404
    try:
        if platform.system() == "Windows":
            os.startfile(str(file_path))
        else:
            subprocess.call(["open" if platform.system() == "Darwin" else "xdg-open", str(file_path)])
        return jsonify({"status": "success", "message": "已打开"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ===== API: 列出输出文件 =====
@app.route("/api/outputs/list")
def list_outputs():
    files = []
    for f in sorted(OUTPUT_DIR.glob("*"), key=lambda p: p.stat().st_mtime, reverse=True):
        files.append({
            "name": f.name,
            "size": f.stat().st_size,
            "time": f.stat().st_mtime,
            "path": str(f)
        })
    return jsonify({"files": files})


# ===== API: 日志 =====
@app.route("/api/logs/recent")
def recent_logs():
    n = request.args.get("n", 50, type=int)
    return jsonify({"logs": logger.get_recent_logs(n)})


@app.route("/api/logs/list")
def list_log_files():
    files = []
    for f in sorted(LOGS_DIR.glob("*.log"), key=lambda p: p.stat().st_mtime, reverse=True):
        files.append({"name": f.name, "size": f.stat().st_size, "time": f.stat().st_mtime})
    return jsonify({"files": files})


# ===== API: 维护接口 - 轻油销量 =====
@app.route("/api/maintenance/light-oil", methods=["GET", "POST"])
def maintenance_light_oil():
    config = load_config()
    if request.method == "GET":
        from base_data import get_light_oil_sales
        return jsonify({"light_oil_sales": get_light_oil_sales(config)})
    elif request.method == "POST":
        data = request.json
        if "file" in request.files:
            # 文件上传模式（预留接口，后续编写具体解析规则）
            f = request.files["file"]
            temp_path = UPLOADS_DIR / f"light_oil_{get_today_str()}.xlsx"
            f.save(str(temp_path))
            return jsonify({"status": "success", "message": "文件已接收，解析规则待开发",
                            "file": str(temp_path),
                            "note": "请后续提供文件样例后编写解析规则"})
        else:
            # 手动输入模式
            station_sales = data.get("station_sales", {})
            config = update_light_oil_sales(config, station_sales)
            return jsonify({"status": "success", "updated": len(station_sales)})


# ===== API: 维护接口 - 月度目标 =====
@app.route("/api/maintenance/targets", methods=["GET", "POST"])
def maintenance_targets():
    config = load_config()
    if request.method == "GET":
        from base_data import get_monthly_targets, DEFAULT_MONTHLY_TARGETS
        configured = config.get("monthly_targets", {})
        if not configured:
            # 未配置时展示内置默认目标
            return jsonify({"monthly_targets": DEFAULT_MONTHLY_TARGETS, "is_default": True})
        return jsonify({"monthly_targets": configured, "is_default": False})
    elif request.method == "POST":
        data = request.json
        year = data.get("year")
        month = data.get("month")
        targets = data.get("targets", {})
        config = update_monthly_targets(config, year, month, targets)
        return jsonify({"status": "success", "message": f"已更新{year}年{month}月目标"})


# ===== API: 维护接口 - 站点片区映射 =====
@app.route("/api/maintenance/station-regions", methods=["GET", "POST"])
def maintenance_station_regions():
    config = load_config()
    if request.method == "GET":
        # 配置为空时返回内置基础数据（完整站点清单），保证前端站点选择可用
        from base_data import get_station_regions
        return jsonify({"station_regions": get_station_regions(config)})
    elif request.method == "POST":
        data = request.json
        region_map = data.get("region_map", {})
        config = update_station_regions(config, region_map)
        return jsonify({"status": "success", "updated": len(region_map)})


# ===== API: 维护接口 - 券系数 =====
@app.route("/api/maintenance/quan-coefficients", methods=["GET", "POST"])
def maintenance_quan_coefficients():
    config = load_config()
    if request.method == "GET":
        from base_data import get_quan_coefficients
        return jsonify({"quan_coefficients": get_quan_coefficients(config)})
    elif request.method == "POST":
        data = request.json
        coefficients = data.get("coefficients", {})
        config = update_quan_coefficients(config, coefficients)
        return jsonify({"status": "success", "updated": len(coefficients)})


# ===== API: 维护接口 - 站点清单 =====
@app.route("/api/maintenance/stations", methods=["GET", "POST"])
def maintenance_stations():
    config = load_config()
    if request.method == "GET":
        return jsonify({"stations": config.get("stations", [])})
    elif request.method == "POST":
        data = request.json
        config["stations"] = data.get("stations", [])
        save_config(config)
        return jsonify({"status": "success", "updated": len(config["stations"])})


# ===== API: 版本信息 =====
@app.route("/api/version")
def version_info():
    version_file = BASE_DIR / "version.json"
    with open(version_file, "r", encoding="utf-8") as f:
        return jsonify(json.load(f))


def run_server(port=18080):
    """启动Flask服务器"""
    print(f"非油报表助手后端启动: http://127.0.0.1:{port}")
    app.run(host="127.0.0.1", port=port, debug=False, threaded=True)


if __name__ == "__main__":
    run_server()
